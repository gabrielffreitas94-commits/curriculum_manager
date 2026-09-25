"""Testes de integração para os novos endpoints da Sprint 2 (/api/v1/resumes).

Cobre o catálogo de PromptSkills, ingestão multimodal de vagas (URL e Documento)
e o Copilot Interativo de IA.
"""

from io import BytesIO
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document
from httpx import AsyncClient

from app.core.url_scraper import SSRFProtectionError
from app.ports.auth_port import AuthUser

AUTH_USER = AuthUser(
    uid="firebase_user_sprint2_test",
    email="sprint2_tester@thothcvs.ai",
    full_name="Sprint2 Tester",
)


@pytest.fixture
async def auth_headers(async_client: AsyncClient) -> dict[str, str]:
    """Sincroniza o usuário no banco relacional e retorna cabeçalho de autenticação Bearer."""
    with patch("app.api.v1.deps.auth_adapter.verify_token", return_value=AUTH_USER):
        res = await async_client.post(
            "/api/v1/auth/sync",
            headers={"Authorization": "Bearer token_sprint2"},
        )
        assert res.status_code == 200
    return {"Authorization": "Bearer token_sprint2"}


def _create_simple_docx() -> bytes:
    """Cria um arquivo DOCX simples em memória."""
    doc = Document()
    doc.add_paragraph("Vaga: Engenheiro de Software Sênior")
    doc.add_paragraph("Requisitos mandatórios: Python e FastAPI.")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_list_prompt_skills_endpoint(
    async_client: AsyncClient,
    auth_headers: dict[str, str],
) -> None:
    """Valida a listagem do catálogo de PromptSkills via endpoint REST."""
    with patch("app.api.v1.deps.auth_adapter.verify_token", return_value=AUTH_USER):
        response = await async_client.get(
            "/api/v1/resumes/prompt-skills",
            headers=auth_headers,
        )

    assert response.status_code == 200
    skills = response.json()
    assert len(skills) >= 5
    slugs = [s["slug"] for s in skills]
    assert "google-xyz" in slugs
    assert "star-method" in slugs
    assert "tech-startup" in slugs
    assert "enterprise-gov" in slugs
    assert "executive-leadership" in slugs


@pytest.mark.asyncio
async def test_ingest_job_from_url_success(
    async_client: AsyncClient,
    auth_headers: dict[str, str],
) -> None:
    """Valida ingestão de vaga via URL pública com sucesso."""
    with (
        patch("app.api.v1.deps.auth_adapter.verify_token", return_value=AUTH_USER),
        patch(
            "app.services.job_ingest_service.safe_fetch_url",
            new_callable=AsyncMock,
            return_value="Descrição da vaga de Engenheiro de Dados.",
        ),
    ):
        response = await async_client.post(
            "/api/v1/resumes/ingest/url",
            json={"url": "https://careers.example.com/job/101"},
            headers=auth_headers,
        )

    assert response.status_code == 200
    data = response.json()
    assert data["source_type"] == "url"
    assert "Descrição da vaga de Engenheiro de Dados." in data["job_description"]
    assert data["char_count"] > 0


@pytest.mark.asyncio
async def test_guardrail_ingest_job_from_url_ssrf_blocked(
    async_client: AsyncClient,
    auth_headers: dict[str, str],
) -> None:
    """VETOR DE AMEAÇA: CWE-918 (Server-Side Request Forgery).
    Tentativa de submeter URL apontando para IP privado ou metadados de nuvem.

    COMPORTAMENTO ESPERADO (FAIL-CLOSED):
    Retornar HTTP 400 com mensagem explícita de bloqueio de segurança.

    RISCO DE REGRESSÃO SILENCIOSA:
    Tratar falhas de SSRF como erros 500 ou não capturar exceções especializadas.

    PREMISSA DO GUARDRAIL (ORÁCULO ABSOLUTO):
    SSRFProtectionError deve ser mapeado para HTTP 400 Bad Request.
    """
    with (
        patch("app.api.v1.deps.auth_adapter.verify_token", return_value=AUTH_USER),
        patch(
            "app.services.job_ingest_service.safe_fetch_url",
            new_callable=AsyncMock,
            side_effect=SSRFProtectionError("IP privado 127.0.0.1 bloqueado."),
        ),
    ):
        response = await async_client.post(
            "/api/v1/resumes/ingest/url",
            json={"url": "http://127.0.0.1:8000/internal-admin"},
            headers=auth_headers,
        )

    assert response.status_code == 400
    assert "Acesso bloqueado por segurança" in response.json()["detail"]


@pytest.mark.asyncio
async def test_ingest_job_from_document_success(
    async_client: AsyncClient,
    auth_headers: dict[str, str],
) -> None:
    """Valida ingestão de vaga via upload de documento DOCX válido."""
    docx_bytes = _create_simple_docx()
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    files = {"file": ("vaga.docx", docx_bytes, mime)}

    with patch("app.api.v1.deps.auth_adapter.verify_token", return_value=AUTH_USER):
        response = await async_client.post(
            "/api/v1/resumes/ingest/document",
            files=files,
            headers=auth_headers,
        )

    assert response.status_code == 200
    data = response.json()
    assert data["source_type"] == "document"
    assert "Engenheiro de Software Sênior" in data["job_description"]


@pytest.mark.asyncio
async def test_copilot_chat_endpoint_success(
    async_client: AsyncClient,
    auth_headers: dict[str, str],
) -> None:
    """Valida endpoint conversacional do Copilot de IA."""
    with (
        patch("app.api.v1.deps.auth_adapter.verify_token", return_value=AUTH_USER),
        patch(
            "app.adapters.gemini_ai_adapter.GeminiAIAdapter.chat_tailoring",
            new_callable=AsyncMock,
        ) as mock_chat,
    ):
        mock_chat.return_value = "Sugiro enfatizar a Fórmula Google XYZ no projeto."
        response = await async_client.post(
            "/api/v1/resumes/copilot/chat",
            json={
                "job_description": "Vaga Tech Lead Python na Startup X.",
                "prompt_skill_slug": "google-xyz",
                "messages": [{"role": "user", "content": "Como otimizar meu CV?"}],
            },
            headers=auth_headers,
        )

    assert response.status_code == 200
    data = response.json()
    assert "Fórmula Google XYZ" in data["reply"]
