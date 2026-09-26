"""Testes unitários para o GeminiResumeParserAdapter.

Valida a extração de currículos em PDF e DOCX com isolamento mock do SDK Gemini,
verificação de encaminhamento de parâmetros e tratamento de exceções de parsing.
"""

from io import BytesIO
from unittest.mock import MagicMock, patch

import docx
import pytest

from app.adapters.gemini_resume_parser_adapter import GeminiResumeParserAdapter
from app.ports.resume_parser_port import (
    CorruptedFileError,
    ParsedProfileDTO,
    ParsingExecutionError,
    ResumeParserError,
)


def _create_mock_parsed_profile_json() -> str:
    """Gera um JSON válido serializado do ParsedProfileDTO para uso nos mocks."""
    dto = ParsedProfileDTO(
        personal_data={
            "full_name": "Gabriel Freitas",
            "headline": "Senior Software Architect",
            "email": "gabriel@example.com",
            "phone": "+55 11 99999-9999",
            "location": "São Paulo, Brasil",
            "linkedin_url": "https://linkedin.com/in/gabrielfreitas",
            "github_url": "https://github.com/gabrielfreitas",
        },
        experiences=[
            {
                "company_name": "Tech Corp",
                "position_title": "Lead Software Engineer",
                "location": "Remoto",
                "work_model": "remote",
                "start_date": "2022-01-01",
                "end_date": None,
                "is_current": True,
                "description": "Liderança técnica de microsserviços.",
                "tech_stack": ["Python", "FastAPI", "Kubernetes"],
                "quantifiable_results": ["Redução de 40% na latência"],
            }
        ],
        educations=[
            {
                "institution_name": "Universidade Tech",
                "degree": "Bacharelado",
                "field_of_study": "Ciência da Computação",
                "start_date": "2016-01-01",
                "end_date": "2020-12-31",
                "is_current": False,
            }
        ],
        skills=[
            {
                "name": "Python",
                "category": "backend",
                "proficiency_level": "advanced",
                "years_of_experience": 6,
            }
        ],
        languages=[
            {
                "language_name": "Inglês",
                "proficiency_level": "fluent",
            }
        ],
    )
    return dto.model_dump_json()


def test_guardrail_missing_api_key_fails_closed() -> None:
    """Valida que a ausência de chave de API recusa a extração imediatamente (Fail-Closed).

    VETOR DE AMEAÇA:
    - CWE-306: Missing Authentication for Critical Function.
    - Impacto Potencial: Inicialização com credenciais nulas causando requisições
      inválidas ou vazamentos.

    COMPORTAMENTO ESPERADO (FAIL-CLOSED):
    - O adaptador DEVE recusar sumariamente a operação levantando ResumeParserError
      se api_key for ausente ou vazia.

    RISCO DE REGRESSÃO SILENCIOSA (ALERTA PARA REFACTOR HUMANO E IA/LLM):
    - Permitir chamadas sem verificar chave ativa geraria exceções não tratadas de terceiros.

    PREMISSA DO GUARDRAIL (ORÁCULO ABSOLUTO):
    - Invocar parse_resume em adaptador sem API Key DEVE falhar com ResumeParserError.
    """
    adapter = GeminiResumeParserAdapter(api_key=None)
    with pytest.raises(ResumeParserError, match="não configurada"):
        import asyncio

        asyncio.run(
            adapter.parse_resume(
                file_bytes=b"%PDF-1.4", mime_type="application/pdf", filename="cv.pdf"
            )
        )


@patch("app.adapters.gemini_resume_parser_adapter.genai.Client")
@pytest.mark.asyncio
async def test_parse_pdf_success(mock_client_class: MagicMock) -> None:
    """Valida o fluxo feliz de extração de PDF com verificação de structured outputs."""
    mock_client = MagicMock()
    mock_client_class.return_value = mock_client

    mock_response = MagicMock()
    mock_response.text = _create_mock_parsed_profile_json()
    mock_client.models.generate_content.return_value = mock_response

    adapter = GeminiResumeParserAdapter(api_key="fake_api_key")
    result = await adapter.parse_resume(
        file_bytes=b"%PDF-1.4 mock content",
        mime_type="application/pdf",
        filename="meu_curriculo.pdf",
    )

    assert isinstance(result, ParsedProfileDTO)
    assert result.personal_data.full_name == "Gabriel Freitas"
    assert len(result.experiences) == 1
    assert result.experiences[0].company_name == "Tech Corp"
    assert result.experiences[0].is_current is True

    # Valida encaminhamento de parâmetros seguros (temperatura 0.0)
    mock_client.models.generate_content.assert_called_once()
    call_kwargs = mock_client.models.generate_content.call_args.kwargs
    assert call_kwargs["model"] == "gemini-2.5-flash"
    assert call_kwargs["config"].temperature == 0.0
    assert call_kwargs["config"].response_schema == ParsedProfileDTO


@patch("app.adapters.gemini_resume_parser_adapter.genai.Client")
@pytest.mark.asyncio
async def test_parse_docx_success(mock_client_class: MagicMock) -> None:
    """Valida o fluxo feliz de extração de texto de DOCX e estruturação via Gemini."""
    mock_client = MagicMock()
    mock_client_class.return_value = mock_client

    mock_response = MagicMock()
    mock_response.text = _create_mock_parsed_profile_json()
    mock_client.models.generate_content.return_value = mock_response

    # Criação de um DOCX válido em memória com parágrafos e tabelas
    doc = docx.Document()
    doc.add_paragraph("Gabriel Freitas - Senior Software Architect")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Tech Corp"
    table.rows[0].cells[1].text = "Lead Software Engineer"
    docx_io = BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    adapter = GeminiResumeParserAdapter(api_key="fake_api_key")
    result = await adapter.parse_resume(
        file_bytes=docx_bytes,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="curriculo.docx",
    )

    assert isinstance(result, ParsedProfileDTO)
    assert result.personal_data.full_name == "Gabriel Freitas"
    mock_client.models.generate_content.assert_called_once()


@pytest.mark.asyncio
async def test_parse_docx_corrupted_bytes_raises_error() -> None:
    """Valida que bytes corrompidos de DOCX levantam CorruptedFileError com tratamento gracioso."""
    adapter = GeminiResumeParserAdapter(api_key="fake_api_key")
    corrupted_bytes = b"PK\x03\x04corrupted zip payload"
    with pytest.raises(CorruptedFileError, match="Falha ao ler estrutura"):
        await adapter.parse_resume(
            file_bytes=corrupted_bytes,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="quebrado.docx",
        )


@pytest.mark.asyncio
async def test_parse_docx_empty_text_raises_error() -> None:
    """Valida que arquivo DOCX sem texto extraível levanta CorruptedFileError."""
    doc = docx.Document()
    docx_io = BytesIO()
    doc.save(docx_io)
    empty_docx_bytes = docx_io.getvalue()

    adapter = GeminiResumeParserAdapter(api_key="fake_api_key")
    with pytest.raises(CorruptedFileError, match="não contém texto extraível"):
        await adapter.parse_resume(
            file_bytes=empty_docx_bytes,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="vazio.docx",
        )


@patch("app.adapters.gemini_resume_parser_adapter.genai.Client")
@pytest.mark.asyncio
async def test_parse_empty_gemini_response_raises_error(mock_client_class: MagicMock) -> None:
    """Valida que retorno vazio do modelo de IA levanta ParsingExecutionError."""
    mock_client = MagicMock()
    mock_client_class.return_value = mock_client

    mock_response = MagicMock()
    mock_response.text = ""  # Resposta vazia
    mock_client.models.generate_content.return_value = mock_response

    adapter = GeminiResumeParserAdapter(api_key="fake_api_key")
    with pytest.raises(ParsingExecutionError, match="resposta vazia"):
        await adapter.parse_resume(
            file_bytes=b"%PDF-1.4 mock",
            mime_type="application/pdf",
            filename="cv.pdf",
        )


@patch("app.adapters.gemini_resume_parser_adapter.genai.Client")
@pytest.mark.asyncio
async def test_parse_gemini_exception_handled(mock_client_class: MagicMock) -> None:
    """Valida que erros da API Gemini são convertidos em ParsingExecutionError tipado."""
    mock_client = MagicMock()
    mock_client_class.return_value = mock_client
    mock_client.models.generate_content.side_effect = Exception("Google Cloud Quota Exceeded")

    adapter = GeminiResumeParserAdapter(api_key="fake_api_key")
    with pytest.raises(ParsingExecutionError, match="Falha ao extrair dados"):
        await adapter.parse_resume(
            file_bytes=b"%PDF-1.4 mock",
            mime_type="application/pdf",
            filename="cv.pdf",
        )
