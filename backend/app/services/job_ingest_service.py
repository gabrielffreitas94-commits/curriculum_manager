"""Serviço de ingestão multimodal de descrições de vagas de emprego.

Suporta ingestão a partir de:
1. Texto bruto copiado e colado.
2. Upload de documentos (.pdf e .docx) com validação de magic bytes (CWE-434).
3. URL pública da oportunidade com blindagem rigorosa contra SSRF (CWE-918).
"""

import time
from io import BytesIO

from docx import Document
from fastapi import HTTPException, status
from pypdf import PdfReader

from app.core.file_security import FileSecurityError, validate_uploaded_file
from app.core.logging import get_logger
from app.core.url_scraper import (
    PayloadTooLargeError,
    ScraperTimeoutError,
    SSRFProtectionError,
    UrlScraperError,
    safe_fetch_url,
)

logger = get_logger(__name__)


class JobIngestService:
    """Orquestrador da ingestão e extração de texto de vagas de emprego."""

    def extract_text_from_document(self, file_bytes: bytes, filename: str) -> str:
        """Valida e extrai o conteúdo textual de arquivos PDF ou DOCX.

        Args:
            file_bytes: Conteúdo binário bruto do arquivo enviado.
            filename: Nome original do arquivo.

        Returns:
            str: Texto extraído do documento.

        Raises:
            HTTPException: 400 se o arquivo for inválido, corrompido ou sem texto extraível;
                           413 se exceder o limite de 5 MB.
        """
        start_time = time.perf_counter()

        try:
            mime_type, ext = validate_uploaded_file(
                file_bytes=file_bytes,
                filename=filename,
                max_size_bytes=5_242_880,  # 5 MB
            )
        except FileSecurityError as exc:
            logger.warning(
                "job_document_upload_rejected",
                filename=filename,
                error=str(exc),
            )
            status_code = (
                status.HTTP_413_CONTENT_TOO_LARGE
                if "5 MB" in str(exc)
                else status.HTTP_400_BAD_REQUEST
            )
            raise HTTPException(status_code=status_code, detail=str(exc)) from exc

        extracted_text = ""
        try:
            if ext == ".pdf":
                reader = PdfReader(BytesIO(file_bytes))
                page_texts = [page.extract_text() or "" for page in reader.pages]
                extracted_text = "\n".join(t.strip() for t in page_texts if t.strip())
            else:
                doc = Document(BytesIO(file_bytes))
                chunks: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if row_cells:
                            chunks.append(" | ".join(row_cells))
                extracted_text = "\n".join(chunks)

        except Exception as exc:
            logger.error(
                "job_document_parsing_failed",
                filename=filename,
                mime_type=mime_type,
                error=str(exc),
                exc_info=True,
            )
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Não foi possível processar o documento enviado: {exc}",
            ) from exc

        if not extracted_text.strip():
            logger.warning("job_document_empty_text", filename=filename)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="O documento enviado não contém texto legível ou extraível.",
            )

        duration_ms = round((time.perf_counter() - start_time) * 1000, 2)
        logger.info(
            "job_document_ingested",
            filename=filename,
            text_length=len(extracted_text),
            duration_ms=duration_ms,
        )
        return extracted_text.strip()

    async def extract_text_from_url(self, url: str) -> str:
        """Efetua o scraping seguro de uma URL de vaga com proteção contra SSRF e DoS.

        Args:
            url: URL informada pelo candidato.

        Returns:
            str: Texto limpo extraído do anúncio da vaga.

        Raises:
            HTTPException: 400 para tentativas de SSRF ou falha de transporte;
                           413 se o payload for excessivo;
                           504 se ocorrer timeout.
        """
        start_time = time.perf_counter()
        try:
            text = await safe_fetch_url(url=url)
        except SSRFProtectionError as exc:
            logger.warning("job_url_ssrf_blocked", url=url, error=str(exc))
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Acesso bloqueado por segurança: {exc}",
            ) from exc
        except PayloadTooLargeError as exc:
            logger.warning("job_url_payload_too_large", url=url, error=str(exc))
            raise HTTPException(
                status_code=status.HTTP_413_CONTENT_TOO_LARGE,
                detail="O conteúdo da página web da vaga excede o limite máximo permitido (2 MB).",
            ) from exc
        except ScraperTimeoutError as exc:
            logger.warning("job_url_timeout", url=url, error=str(exc))
            raise HTTPException(
                status_code=status.HTTP_504_GATEWAY_TIMEOUT,
                detail="Tempo limite esgotado ao tentar acessar a URL da vaga.",
            ) from exc
        except UrlScraperError as exc:
            logger.error("job_url_scraper_error", url=url, error=str(exc))
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Não foi possível obter a vaga a partir da URL informada: {exc}",
            ) from exc

        if not text.strip():
            logger.warning("job_url_empty_text", url=url)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Nenhum conteúdo textual legível foi encontrado na página da vaga.",
            )

        duration_ms = round((time.perf_counter() - start_time) * 1000, 2)
        logger.info(
            "job_url_ingested",
            url=url,
            text_length=len(text),
            duration_ms=duration_ms,
        )
        return text.strip()
